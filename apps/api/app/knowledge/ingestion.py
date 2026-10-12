import io
import re
import zipfile
from pathlib import Path

from docx import Document
from pypdf import PdfReader

MAX_DOCUMENT_BYTES = 5_000_000
MAX_EXTRACTED_CHARS = 100_000
MAX_PDF_PAGES = 200
MAX_DOCX_FILES = 1_000
MAX_DOCX_EXPANDED_BYTES = 20_000_000
ALLOWED_SUFFIXES = frozenset({".txt", ".md", ".pdf", ".docx"})


class DocumentIngestionError(ValueError):
    pass


def extract_document(filename: str, content: bytes) -> tuple[str, str]:
    suffix = Path(filename).suffix.casefold()
    if suffix not in ALLOWED_SUFFIXES:
        raise DocumentIngestionError("Only TXT, Markdown, PDF, and DOCX documents are supported")
    if not content:
        raise DocumentIngestionError("Document is empty")
    if len(content) > MAX_DOCUMENT_BYTES:
        raise DocumentIngestionError("Document exceeds 5 MB")
    if suffix in {".txt", ".md"}:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as error:
            raise DocumentIngestionError("Text documents must use UTF-8") from error
    elif suffix == ".pdf":
        text = _pdf_text(content)
    else:
        text = _docx_text(content)
    cleaned = _clean(text)
    if not cleaned:
        raise DocumentIngestionError("Document contains no extractable text")
    return cleaned, suffix.removeprefix(".")


def _pdf_text(content: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(content), strict=True)
        if reader.is_encrypted:
            raise DocumentIngestionError("Encrypted PDF documents are not supported")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise DocumentIngestionError("PDF exceeds 200 pages")
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    except DocumentIngestionError:
        raise
    except Exception as error:
        raise DocumentIngestionError("PDF could not be parsed safely") from error


def _docx_text(content: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_DOCX_FILES:
                raise DocumentIngestionError("DOCX contains too many embedded files")
            if any(entry.flag_bits & 0x1 for entry in entries):
                raise DocumentIngestionError("Encrypted DOCX documents are not supported")
            if sum(entry.file_size for entry in entries) > MAX_DOCX_EXPANDED_BYTES:
                raise DocumentIngestionError("DOCX expanded content exceeds 20 MB")
        document = Document(io.BytesIO(content))
        values = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            values.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(values)
    except DocumentIngestionError:
        raise
    except (ValueError, KeyError, zipfile.BadZipFile) as error:
        raise DocumentIngestionError("DOCX could not be parsed safely") from error


def _clean(value: str) -> str:
    value = value.replace("\x00", "")
    value = re.sub(r"[\x01-\x08\x0b\x0c\x0e-\x1f\x7f]", "", value)
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n{3,}", "\n\n", value).strip()
    if len(value) > MAX_EXTRACTED_CHARS:
        raise DocumentIngestionError("Extracted text exceeds 100,000 characters")
    return value
